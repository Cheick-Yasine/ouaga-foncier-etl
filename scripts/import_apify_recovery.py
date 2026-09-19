"""Importe en masse des exports Apify avec reprise sûre par lots.

Le script :
1. lit tous les .xlsx/.xlsm/.csv/.json d'un dossier ;
2. normalise les formats Apify ;
3. déduplique uniquement les doublons techniques ;
4. conserve le filtrage métier existant ;
5. structure les candidats avec OpenAI par petits lots ;
6. sauvegarde les résultats LLM après chaque lot ;
7. écrit immédiatement les annonces valides dans Neon ;
8. reprend sans repayer les lots déjà traités après une coupure.

Aucun filtre de date n'est appliqué.

Exemples :
    python scripts/import_apify_recovery.py data/apify_recovery --prepare-only
    python scripts/import_apify_recovery.py data/apify_recovery
    python scripts/import_apify_recovery.py data/apify_recovery --batch-size 50 --concurrency 10
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import hashlib
import json
import math
import re
import sys
import time
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

# IMPORTANT : charger .env avant config, car config.DATABASE_URL est lu à l'import.
load_dotenv(ROOT / ".env")

import config
import processor

logger = config.configurer_logging()

EXTENSIONS_ACCEPTEES = {".xlsx", ".xlsm", ".csv", ".json"}
CACHE_LLM_PATH = config.STATE_DIR / "apify_recovery_llm_cache.json"
CHECKPOINT_PATH = config.STATE_DIR / "apify_recovery_checkpoint.json"
BIGINT_MAX = 9_223_372_036_854_775_807
BIGINT_MIN = -9_223_372_036_854_775_808

# Les exports Apify varient selon l'Actor. On compare des noms de colonnes
# normalisés (minuscules, sans accents ni ponctuation).
ALIASES = {
    "id": {
        "id", "postid", "postidentifier", "facebookid", "facebookpostid",
        "postfacebookid", "postpk", "postkey",
    },
    "texte": {
        "text", "texte", "message", "content", "contenu", "posttext",
        "postmessage", "body", "description", "caption",
    },
    "url": {
        "url", "posturl", "facebookurl", "permalink", "postlink",
        "link", "postpermalink",
    },
    "date_publication": {
        "time", "date", "datetime", "timestamp", "createdat", "createdtime",
        "publishedat", "publicationdate", "datepublication", "postdate",
        "posttime",
    },
    "groupe_nom": {
        "groupname", "grouptitle", "facebookgroupname", "groupe", "group",
    },
}


def _charger_groupes_connus() -> dict[str, str]:
    """Charge la correspondance ID Facebook -> nom du groupe suivi."""
    chemin = ROOT / "groups.csv"
    if not chemin.exists():
        return {}

    groupes: dict[str, str] = {}
    with chemin.open("r", encoding="utf-8-sig", newline="") as f:
        for ligne in csv.DictReader(f):
            group_id = str(ligne.get("id") or "").strip()
            nom = str(ligne.get("nom") or "").strip()
            if group_id and nom:
                groupes[group_id] = nom
    return groupes


_GROUPES_PAR_ID = _charger_groupes_connus()
_RE_GROUP_ID_URL = re.compile(r"/groups/(\d+)(?:/|$)", re.I)


def _nom_groupe_depuis_url(url: str | None) -> str | None:
    """Privilégie l'ID du groupe dans l'URL, plus fiable que l'auteur du post."""
    if not url:
        return None
    match = _RE_GROUP_ID_URL.search(str(url))
    if not match:
        return None
    return _GROUPES_PAR_ID.get(match.group(1))


def _normaliser_nom_colonne(valeur: Any) -> str:
    texte = unicodedata.normalize("NFKD", str(valeur or ""))
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", texte.lower())


def _nettoyer_cellule(valeur: Any) -> Any:
    if isinstance(valeur, (datetime, date)):
        return valeur.isoformat()
    if valeur is None:
        return None
    if isinstance(valeur, str):
        valeur = valeur.strip()
        return valeur or None
    return valeur


def _aplatir_dict(obj: dict[str, Any], prefixe: str = "") -> dict[str, Any]:
    sortie: dict[str, Any] = {}
    for cle, valeur in obj.items():
        nom = f"{prefixe}{cle}" if prefixe else str(cle)
        if isinstance(valeur, dict):
            sortie.update(_aplatir_dict(valeur, nom))
        else:
            sortie[nom] = valeur
    return sortie


def _valeur_alias(ligne: dict[str, Any], champ: str) -> Any:
    index = {
        _normaliser_nom_colonne(cle): _nettoyer_cellule(valeur)
        for cle, valeur in ligne.items()
    }
    for alias in ALIASES[champ]:
        if alias in index and index[alias] not in (None, ""):
            return index[alias]
    return None


_RE_ID_URL = [
    re.compile(r"/posts/(\d+)", re.I),
    re.compile(r"/permalink/(\d+)", re.I),
    re.compile(r"[?&]story_fbid=(\d+)", re.I),
    re.compile(r"[?&]fbid=(\d+)", re.I),
]


def _extraire_id_url(url: str | None) -> str | None:
    if not url:
        return None
    for motif in _RE_ID_URL:
        match = motif.search(str(url))
        if match:
            return match.group(1)
    return None


def _normaliser_id(valeur: Any) -> str | None:
    if valeur is None:
        return None
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    texte = str(valeur).strip()
    if not texte:
        return None
    if re.fullmatch(r"\d+\.0", texte):
        return texte[:-2]
    return texte


def _id_synthetique(
    url: str | None,
    texte: str | None,
    date_pub: str | None,
    source: str,
) -> str:
    base = "|".join([
        str(url or "").strip(),
        str(texte or "").strip(),
        str(date_pub or "").strip(),
        source,
    ])
    digest = hashlib.sha256(base.encode("utf-8")).hexdigest()[:32]
    return f"apify_{digest}"


def _normaliser_ligne(ligne: dict[str, Any], source: str) -> dict[str, Any]:
    ligne = _aplatir_dict(ligne)

    texte = _valeur_alias(ligne, "texte")
    url = _valeur_alias(ligne, "url")
    date_pub = _valeur_alias(ligne, "date_publication")
    groupe_nom = (
        _nom_groupe_depuis_url(str(url) if url else None)
        or _valeur_alias(ligne, "groupe_nom")
        or source
    )

    post_id = _normaliser_id(_valeur_alias(ligne, "id"))
    if not post_id:
        post_id = _extraire_id_url(str(url) if url else None)
    if not post_id:
        post_id = _id_synthetique(
            str(url) if url else None,
            str(texte) if texte else None,
            str(date_pub) if date_pub else None,
            source,
        )

    return {
        "id": post_id,
        "groupe_nom": str(groupe_nom).strip(),
        "url": str(url).strip() if url else "",
        "date_publication": str(date_pub).strip() if date_pub else "",
        "date_incertaine": not bool(date_pub),
        "texte": str(texte).strip() if texte else "",
    }


def _lire_excel(chemin: Path) -> Iterable[dict[str, Any]]:
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    try:
        for feuille in classeur.worksheets:
            lignes = feuille.iter_rows(values_only=True)
            try:
                entetes_brutes = next(lignes)
            except StopIteration:
                continue

            entetes = [
                str(v).strip() if v is not None else f"colonne_{i}"
                for i, v in enumerate(entetes_brutes, start=1)
            ]
            for valeurs in lignes:
                if not any(v not in (None, "") for v in valeurs):
                    continue
                yield {
                    entetes[i]: valeurs[i] if i < len(valeurs) else None
                    for i in range(len(entetes))
                }
    finally:
        classeur.close()


def _detecter_dialecte_csv(chemin: Path) -> csv.Dialect:
    echantillon = chemin.read_text(encoding="utf-8-sig", errors="replace")[:8192]
    try:
        return csv.Sniffer().sniff(echantillon, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def _lire_csv(chemin: Path) -> Iterable[dict[str, Any]]:
    dialecte = _detecter_dialecte_csv(chemin)
    with chemin.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        for ligne in csv.DictReader(f, dialect=dialecte):
            yield dict(ligne)


def _lire_json(chemin: Path) -> Iterable[dict[str, Any]]:
    with chemin.open("r", encoding="utf-8") as f:
        contenu = json.load(f)

    if isinstance(contenu, list):
        objets = contenu
    elif isinstance(contenu, dict):
        objets = None
        for cle in ("items", "data", "results", "posts"):
            valeur = contenu.get(cle)
            if isinstance(valeur, list):
                objets = valeur
                break
        if objets is None:
            objets = [contenu]
    else:
        objets = []

    for objet in objets:
        if isinstance(objet, dict):
            yield objet


def _lire_fichier(chemin: Path) -> Iterable[dict[str, Any]]:
    suffixe = chemin.suffix.lower()
    if suffixe in {".xlsx", ".xlsm"}:
        yield from _lire_excel(chemin)
    elif suffixe == ".csv":
        yield from _lire_csv(chemin)
    elif suffixe == ".json":
        yield from _lire_json(chemin)


def _cle_dedup(post: dict[str, Any]) -> str:
    post_id = str(post.get("id") or "").strip()
    if post_id and not post_id.startswith("apify_"):
        return f"id:{post_id}"

    url = str(post.get("url") or "").strip()
    if url:
        return f"url:{url}"

    return f"id:{post_id}"


def preparer_import(dossier: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    fichiers = sorted(
        p for p in dossier.rglob("*")
        if p.is_file() and p.suffix.lower() in EXTENSIONS_ACCEPTEES
    )
    if not fichiers:
        raise FileNotFoundError(
            f"Aucun fichier .xlsx/.xlsm/.csv/.json trouvé dans {dossier}"
        )

    uniques: dict[str, dict[str, Any]] = {}
    fichiers_ok = 0
    fichiers_erreur: list[tuple[str, str]] = []
    nb_lignes = 0
    nb_doublons = 0
    nb_ids_synthetiques = 0

    for chemin in fichiers:
        source = chemin.stem
        compteur_fichier = 0
        try:
            for ligne in _lire_fichier(chemin):
                nb_lignes += 1
                compteur_fichier += 1
                post = _normaliser_ligne(ligne, source=source)
                if post["id"].startswith("apify_"):
                    nb_ids_synthetiques += 1

                cle = _cle_dedup(post)
                if cle in uniques:
                    nb_doublons += 1
                uniques[cle] = post

            fichiers_ok += 1
            print(f"OK  {chemin.name}: {compteur_fichier} ligne(s)")
        except Exception as exc:
            fichiers_erreur.append((chemin.name, str(exc)))
            print(f"ERREUR  {chemin.name}: {exc}")

    stats = {
        "fichiers_trouves": len(fichiers),
        "fichiers_ok": fichiers_ok,
        "fichiers_erreur": fichiers_erreur,
        "lignes_brutes": nb_lignes,
        "doublons": nb_doublons,
        "uniques": len(uniques),
        "ids_synthetiques": nb_ids_synthetiques,
    }
    return list(uniques.values()), stats


def _empreinte_candidats(candidats: list[dict[str, Any]]) -> str:
    h = hashlib.sha256()
    for post in sorted(candidats, key=lambda p: str(p.get("id") or "")):
        h.update(str(post.get("id") or "").encode("utf-8"))
        h.update(b"\0")
        h.update(str(post.get("texte_nettoye") or "").encode("utf-8"))
        h.update(b"\n")
    return h.hexdigest()


def _ecrire_json_atomique(chemin: Path, contenu: dict[str, Any]) -> None:
    chemin.parent.mkdir(parents=True, exist_ok=True)
    temp = chemin.with_suffix(chemin.suffix + ".tmp")
    temp.write_text(
        json.dumps(contenu, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    temp.replace(chemin)


def _charger_etat(empreinte: str) -> tuple[dict[str, Any], dict[str, Any]]:
    cache = {"empreinte": empreinte, "results": {}}
    checkpoint = {"empreinte": empreinte, "db_done_ids": []}

    if CACHE_LLM_PATH.exists():
        try:
            contenu = json.loads(CACHE_LLM_PATH.read_text(encoding="utf-8"))
            if contenu.get("empreinte") == empreinte:
                cache = contenu
        except (OSError, json.JSONDecodeError):
            pass

    if CHECKPOINT_PATH.exists():
        try:
            contenu = json.loads(CHECKPOINT_PATH.read_text(encoding="utf-8"))
            if contenu.get("empreinte") == empreinte:
                checkpoint = contenu
        except (OSError, json.JSONDecodeError):
            pass

    return cache, checkpoint


def _sauvegarder_etat(cache: dict[str, Any], checkpoint: dict[str, Any]) -> None:
    cache["updated_at"] = datetime.now(timezone.utc).isoformat()
    checkpoint["updated_at"] = datetime.now(timezone.utc).isoformat()
    _ecrire_json_atomique(CACHE_LLM_PATH, cache)
    _ecrire_json_atomique(CHECKPOINT_PATH, checkpoint)


def _preparer_schema_neon() -> None:
    """Crée le schéma si besoin et agrandit les colonnes numériques existantes."""
    if not config.DATABASE_URL:
        raise ValueError("DATABASE_URL absente. Vérifiez votre fichier .env.")

    with psycopg.connect(config.DATABASE_URL) as conn:
        with conn.cursor() as cur:
            cur.execute(processor.SCHEMA_SQL)
            # INTEGER s'arrête à 2 147 483 647 : insuffisant pour certains
            # prix immobiliers. BIGINT évite l'erreur NumericValueOutOfRange.
            cur.execute(
                "ALTER TABLE annonces "
                "ALTER COLUMN prix_fcfa TYPE BIGINT USING prix_fcfa::BIGINT"
            )
            cur.execute(
                "ALTER TABLE annonces "
                "ALTER COLUMN superficie_m2 TYPE BIGINT USING superficie_m2::BIGINT"
            )
        conn.commit()


def _securiser_grands_entiers(annonces: list[dict[str, Any]]) -> None:
    """Évite qu'une valeur LLM manifestement hors plage BIGINT casse un lot."""
    for annonce in annonces:
        for champ in ("prix_fcfa", "superficie_m2"):
            valeur = annonce.get(champ)
            if isinstance(valeur, int) and not (BIGINT_MIN <= valeur <= BIGINT_MAX):
                logger.warning(
                    "%s hors plage BIGINT pour post %s (%s) -> null",
                    champ,
                    annonce.get("id"),
                    valeur,
                )
                annonce[champ] = None


def _lots(elements: list[Any], taille: int) -> Iterable[list[Any]]:
    for i in range(0, len(elements), taille):
        yield elements[i:i + taille]


async def _traiter_reprise(
    posts: list[dict[str, Any]],
    *,
    batch_size: int,
    concurrency: int,
    export_master_xlsx: bool = False,
) -> tuple[int, int, Path, Path | None]:
    candidats, rejetes_niveau1 = processor.filtrer_candidats(posts)
    empreinte = _empreinte_candidats(candidats)
    cache, checkpoint = _charger_etat(empreinte)

    results: dict[str, Any] = cache.setdefault("results", {})
    db_done = set(str(x) for x in checkpoint.setdefault("db_done_ids", []))

    # Le cache LLM doit survivre à un échec DB. On peut donc reprendre ici sans
    # refaire les appels OpenAI d'un lot déjà structuré.
    valides_cachees_a_ecrire = [
        entree["record"]
        for post_id, entree in results.items()
        if entree.get("status") == "valid"
        and post_id not in db_done
        and isinstance(entree.get("record"), dict)
    ]

    if valides_cachees_a_ecrire:
        print(
            f"Reprise : {len(valides_cachees_a_ecrire)} résultat(s) LLM déjà "
            "sauvegardé(s) à écrire dans Neon sans nouvel appel API."
        )
        for lot in _lots(valides_cachees_a_ecrire, batch_size):
            _securiser_grands_entiers(lot)
            processor.upsert_annonces(lot)
            db_done.update(str(a["id"]) for a in lot)
            checkpoint["db_done_ids"] = sorted(db_done)
            _sauvegarder_etat(cache, checkpoint)

    a_traiter = [
        p for p in candidats
        if str(p["id"]) not in results
        or results[str(p["id"])].get("status") == "failed"
    ]

    nb_cached = len(candidats) - len(a_traiter)
    print(
        f"Checkpoint : {nb_cached}/{len(candidats)} candidat(s) déjà structuré(s), "
        f"{len(a_traiter)} restant(s)."
    )

    # Concurrence plus élevée uniquement pour ce rattrapage ; le daily garde
    # sa valeur habituelle dans config.py.
    config.LLM_MAX_CONCURRENCE = max(1, concurrency)

    total_lots = max(1, math.ceil(len(a_traiter) / batch_size)) if a_traiter else 0
    debut_global = time.monotonic()

    for index, lot in enumerate(_lots(a_traiter, batch_size), start=1):
        debut_lot = time.monotonic()
        print(
            f"\nLot {index}/{total_lots} : {len(lot)} candidat(s) -> OpenAI "
            f"(concurrence={config.LLM_MAX_CONCURRENCE})"
        )

        valides, non_valides = await processor.structurer_lot(lot)
        _securiser_grands_entiers(valides)

        # Sauvegarder le résultat LLM AVANT l'écriture DB : même si Neon tombe
        # juste après, ce lot ne sera pas repassé par OpenAI au prochain run.
        for record in valides:
            results[str(record["id"])] = {
                "status": "valid",
                "record": record,
            }

        for record in non_valides:
            post_id = str(record.get("id") or "")
            motif = record.get("motif_rejet")
            # Les vrais refus métier sont définitifs. Un échec API/validation
            # reste retentable au prochain lancement.
            statut = "failed" if motif == "echec_api_ou_validation" else "rejected"
            results[post_id] = {
                "status": statut,
                "record": record,
            }

        _sauvegarder_etat(cache, checkpoint)

        if valides:
            processor.upsert_annonces(valides)
            db_done.update(str(a["id"]) for a in valides)
            checkpoint["db_done_ids"] = sorted(db_done)
            _sauvegarder_etat(cache, checkpoint)

        duree = time.monotonic() - debut_lot
        traites = min(index * batch_size, len(a_traiter))
        restants = max(0, len(a_traiter) - traites)
        vitesse = traites / max(time.monotonic() - debut_global, 0.001)
        eta = restants / vitesse if vitesse > 0 else 0

        print(
            f"Lot {index}/{total_lots} terminé : "
            f"{len(valides)} valide(s), {len(non_valides)} rejetée(s)/échouée(s), "
            f"Neon OK | {duree:.1f}s | ETA ~ {eta / 60:.1f} min"
        )

    # Reconstituer le résultat total depuis le cache, y compris ce qui venait
    # d'un run précédent.
    valides_totales = [
        entree["record"]
        for entree in results.values()
        if entree.get("status") == "valid"
        and isinstance(entree.get("record"), dict)
    ]
    rejetes_niveau2 = [
        entree["record"]
        for entree in results.values()
        if entree.get("status") in {"rejected", "failed"}
        and isinstance(entree.get("record"), dict)
    ]

    horodatage = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")

    print("\nFinalisation 1/3 : enregistrement des statistiques du run...")
    processor.enregistrer_run(
        "recovery",
        len(posts),
        len(candidats),
        len(valides_totales),
    )
    print("Finalisation 1/3 OK.")

    print("Finalisation 2/3 : export CSV des annonces valides...")
    chemin_csv = config.PROCESSED_DIR / f"annonces_{horodatage}.csv"
    processor.exporter_csv(valides_totales, chemin_csv)
    print("Finalisation 2/3 OK.")

    print("Finalisation 3/3 : export audit des rejets...")
    processor.exporter_json_audit(
        rejetes_niveau1 + rejetes_niveau2,
        config.PROCESSED_DIR / f"rejetes_{horodatage}.json",
    )
    print("Finalisation 3/3 OK.")

    # La régénération de l'Excel maître parcourt toute la table Neon et peut
    # être lente/inutile pour un import massif. Elle devient optionnelle.
    chemin_xlsx = None
    if export_master_xlsx:
        print("Export Excel maître demandé : génération en cours...")
        chemin_xlsx = processor.exporter_xlsx_depuis_db()
        print("Export Excel maître OK.")

    return len(candidats), len(valides_totales), chemin_csv, chemin_xlsx


async def _executer(args: argparse.Namespace) -> int:
    dossier = Path(args.dossier).expanduser().resolve()
    if not dossier.exists() or not dossier.is_dir():
        raise FileNotFoundError(f"Dossier introuvable : {dossier}")

    posts, stats = preparer_import(dossier)

    sortie = Path(args.output).expanduser()
    if not sortie.is_absolute():
        sortie = ROOT / sortie
    sortie.parent.mkdir(parents=True, exist_ok=True)
    sortie.write_text(
        json.dumps(posts, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )

    print("\nIMPORT APIFY - PREPARATION")
    print("-" * 42)
    print(f"Fichiers trouvés       : {stats['fichiers_trouves']}")
    print(f"Fichiers lus           : {stats['fichiers_ok']}")
    print(f"Lignes brutes          : {stats['lignes_brutes']}")
    print(f"Doublons techniques    : {stats['doublons']}")
    print(f"Posts uniques          : {stats['uniques']}")
    print(f"IDs synthétiques       : {stats['ids_synthetiques']}")
    print(f"JSON normalisé         : {sortie}")

    if stats["fichiers_erreur"]:
        print("\nFichiers ignorés à cause d'une erreur :")
        for nom, erreur in stats["fichiers_erreur"]:
            print(f"  - {nom}: {erreur}")

    if args.prepare_only:
        print("\nPréparation terminée (--prepare-only) : aucune écriture dans Neon.")
        return 0

    if not posts:
        print("\nAucun post à traiter.")
        return 0

    _preparer_schema_neon()
    print(
        "\nSchéma Neon vérifié : prix_fcfa et superficie_m2 acceptent maintenant "
        "les valeurs BIGINT."
    )
    print(
        "Lancement filtrage -> OpenAI par lots -> sauvegarde checkpoint -> Neon..."
    )

    nb_candidats, nb_valides, chemin_csv, chemin_xlsx = await _traiter_reprise(
        posts,
        batch_size=max(1, args.batch_size),
        concurrency=max(1, args.concurrency),
        export_master_xlsx=args.export_master_xlsx,
    )

    print("\nIMPORT APIFY - RESULTAT")
    print("-" * 42)
    print(f"Posts bruts chargés    : {len(posts)}")
    print(f"Candidats immobiliers  : {nb_candidats}")
    print(f"Annonces valides       : {nb_valides}")
    print(f"CSV du run             : {chemin_csv}")
    print(
        f"Excel maître           : {chemin_xlsx if chemin_xlsx else 'non généré (optionnel)'}"
    )
    print(f"Cache LLM              : {CACHE_LLM_PATH}")
    print(f"Checkpoint Neon        : {CHECKPOINT_PATH}")
    print("\nOK - import terminé. Les lots déjà traités sont réutilisables après coupure.")
    return 0


def parser_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Fusionne les exports Apify puis exécute filtrage + LLM + Neon "
            "par lots avec reprise sûre, sans filtre de date."
        )
    )
    parser.add_argument(
        "dossier",
        nargs="?",
        default=str(ROOT / "data" / "apify_recovery"),
        help="Dossier contenant les exports Apify.",
    )
    parser.add_argument(
        "--output",
        default="data/raw/apify_recovery_merged.json",
        help="JSON brut normalisé produit avant le traitement.",
    )
    parser.add_argument(
        "--prepare-only",
        action="store_true",
        help="Fusionne/normalise/déduplique sans OpenAI ni écriture Neon.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=50,
        help="Nombre de candidats par lot LLM/Neon (défaut : 50).",
    )
    parser.add_argument(
        "--concurrency",
        type=int,
        default=10,
        help="Nombre maximal d'appels OpenAI simultanés pour ce rattrapage (défaut : 10).",
    )
    parser.add_argument(
        "--export-master-xlsx",
        action="store_true",
        help=(
            "Régénère aussi l'Excel maître depuis toute la base Neon. "
            "Désactivé par défaut car cette étape peut être lente."
        ),
    )
    return parser.parse_args()


def main() -> int:
    return asyncio.run(_executer(parser_arguments()))


if __name__ == "__main__":
    raise SystemExit(main())
